"""Full-script interaction regression tests for ui.py's Custom Template
Builder, using streamlit.testing.v1.AppTest to drive the real app instead of
extracted functions.

These bugs are not reachable via unit tests of individual functions: the
defect was in *how the whole script reruns* when a button handler calls
st.rerun() before later widgets in the script have been instantiated for
that pass. Streamlit prunes session_state for any keyed widget not seen
during a script pass, so an early st.rerun() wiped out every already-entered
Service Flow field (song numbers, leaders, verses/chorus selections) on
every Add/Remove/Move click in the Custom Template Builder.
"""
import json
from pathlib import Path

import pytest

pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parent.parent
UI_PY = ROOT / "ui.py"


def start_custom_template_with_songs(count):
    at = AppTest.from_file(str(UI_PY))
    at.run(timeout=30)
    at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
    for _ in range(count):
        at.selectbox(key="custom_item_type_select").set_value("song").run()
        at.button(key="custom_add_item").click().run()
    return at


class TestCustomTemplateDistinctLeaders:
    """Regression for: adding several songs with different leaders collapsed
    every leader field down to whichever name was typed last."""

    def test_each_song_gets_its_own_leader_field(self):
        at = start_custom_template_with_songs(3)
        assert list(at.exception) == []

        leader_inputs = [ti for ti in at.text_input if ti.key and ti.key.startswith("leader_Song Leader")]
        assert len(leader_inputs) == 3
        assert len({ti.key for ti in leader_inputs}) == 3

        for ti, name in zip(leader_inputs, ["Alice", "Bob", "Carol"]):
            at.text_input(key=ti.key).set_value(name).run()

        final = {ti.key: at.text_input(key=ti.key).value for ti in leader_inputs}
        assert final == {
            "leader_Song Leader 1": "Alice",
            "leader_Song Leader 2": "Bob",
            "leader_Song Leader 3": "Carol",
        }


class TestCustomTemplateReorderPreservesData:
    """Regression for: moving an item in the Custom Template Builder reset
    every already-entered Service Flow field back to its default."""

    def test_apply_order_preserves_song_numbers(self):
        at = start_custom_template_with_songs(3)
        num_inputs = [ni for ni in at.number_input if ni.key and ni.key.startswith("song_song-")]
        for ni, val in zip(num_inputs, [12, 67, 275]):
            at.number_input(key=ni.key).set_value(val).run()

        before = {ni.key: at.number_input(key=ni.key).value for ni in num_inputs}

        # Send song-1 (row 1) to position 2, swapping it with song-2.
        at.number_input(key="custom_pos_song-1_song").set_value(2).run()
        at.number_input(key="custom_pos_song-2_song").set_value(1).run()
        at.button(key="custom_apply_order").click().run()
        assert list(at.exception) == []

        after = {ni.key: at.number_input(key=ni.key).value for ni in num_inputs}
        assert after == before

        order = [item.get("id") for item in at.session_state["custom_template_items"]]
        assert order == ["song-2", "song-1", "song-3"]

    def test_apply_order_moves_last_item_to_top_in_one_click(self):
        """Regression for: with the old Up/Down-arrow UI, moving an item
        many rows in one direction required re-locating and re-clicking a
        *different* button after every single-step move (the arrows are
        keyed by row index, not item identity) -- clicking the same visual
        row position twice just swapped the same two items back and forth,
        which looked like items randomly changing identity and made far
        moves painful. Typing a target position once and clicking Apply
        should move an item from the bottom to the top in a single step."""
        at = start_custom_template_with_songs(5)
        num_inputs = [ni for ni in at.number_input if ni.key and ni.key.startswith("song_song-")]
        for ni, val in zip(num_inputs, [11, 22, 33, 44, 55]):
            at.number_input(key=ni.key).set_value(val).run()

        at.number_input(key="custom_pos_song-5_song").set_value(1).run()
        at.button(key="custom_apply_order").click().run()
        assert list(at.exception) == []

        order = [item.get("id") for item in at.session_state["custom_template_items"]]
        assert order == ["song-5", "song-1", "song-2", "song-3", "song-4"]
        assert at.number_input(key="song_song-5").value == 55
        assert at.number_input(key="song_song-1").value == 11

    def test_insert_at_top_places_new_item_first(self):
        at = start_custom_template_with_songs(2)
        at.checkbox(key="custom_add_to_top").set_value(True).run()
        at.selectbox(key="custom_item_type_select").set_value("song").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        order = [item.get("id") for item in at.session_state["custom_template_items"]]
        assert order == ["song-3", "song-1", "song-2"]

    def test_add_to_top_unchecked_still_appends_at_bottom(self):
        at = start_custom_template_with_songs(2)
        at.selectbox(key="custom_item_type_select").set_value("song").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        order = [item.get("id") for item in at.session_state["custom_template_items"]]
        assert order == ["song-1", "song-2", "song-3"]

    def test_add_item_preserves_existing_song_and_leader_data(self):
        at = start_custom_template_with_songs(2)
        num_inputs = [ni for ni in at.number_input if ni.key and ni.key.startswith("song_song-")]
        for ni, val in zip(num_inputs, [12, 67]):
            at.number_input(key=ni.key).set_value(val).run()
        leader_inputs = [ti for ti in at.text_input if ti.key and ti.key.startswith("leader_Song Leader")]
        for ti, name in zip(leader_inputs, ["Alice", "Bob"]):
            at.text_input(key=ti.key).set_value(name).run()

        before_nums = {ni.key: at.number_input(key=ni.key).value for ni in num_inputs}
        before_leaders = {ti.key: at.text_input(key=ti.key).value for ti in leader_inputs}

        at.selectbox(key="custom_item_type_select").set_value("song").run()
        at.button(key="custom_add_item").click().run()

        after_nums = {ni.key: at.number_input(key=ni.key).value for ni in num_inputs}
        after_leaders = {ti.key: at.text_input(key=ti.key).value for ti in leader_inputs}
        assert after_nums == before_nums
        assert after_leaders == before_leaders

    def test_remove_item_preserves_surviving_items_data(self):
        at = start_custom_template_with_songs(3)
        num_inputs = {ni.key: ni for ni in at.number_input if ni.key and ni.key.startswith("song_song-")}
        for key, val in zip(["song_song-1", "song_song-2", "song_song-3"], [12, 67, 275]):
            at.number_input(key=key).set_value(val).run()

        at.button(key="custom_remove_1").click().run()
        assert list(at.exception) == []

        order = [item.get("id") for item in at.session_state["custom_template_items"]]
        assert order == ["song-1", "song-3"]
        assert at.number_input(key="song_song-1").value == 12
        assert at.number_input(key="song_song-3").value == 275


class TestCustomTemplateVerseChorusPersistence:
    """Regression for: deselecting a verse/chorus entry, then triggering an
    unrelated rerun, silently re-selected everything again.

    Uses pftl-738, the one song sample kept tracked in git specifically for
    fixture use (see .gitignore's ehsf/ allowlist) -- it has verses 1-4.
    """

    def test_verse_deselection_survives_unrelated_rerun(self):
        at = start_custom_template_with_songs(1)
        at.number_input(key="song_song-1").set_value(738).run()
        assert list(at.exception) == []

        verses = at.multiselect(key="verses_song-1")
        assert verses.value == [1, 2, 3, 4]

        at.multiselect(key="verses_song-1").set_value([1, 2, 3]).run()
        assert at.multiselect(key="verses_song-1").value == [1, 2, 3]

        at.selectbox(key="custom_item_type_select").set_value("reading").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        assert at.multiselect(key="verses_song-1").value == [1, 2, 3]


def start_normal_template(template_name="wednesday", song_number=738):
    """wednesday is the smallest normal template (3 songs), keeping this fast."""
    at = AppTest.from_file(str(UI_PY))
    at.run(timeout=30)
    at.selectbox(key="template_select").set_value(template_name).run()
    song_inputs = [ni for ni in at.number_input if ni.key and ni.key.startswith("song_song-")]
    for ni in song_inputs:
        at.number_input(key=ni.key).set_value(song_number).run()
    return at


def _generate_button(at):
    return [b for b in at.button if "Generate" in (b.label or "")][0]


def _cleanup_generated_files(pptx_path):
    """Remove exactly what create_worship_files/generate_presentation wrote
    for one test run, without touching any sibling date folders."""
    pptx_path = Path(pptx_path)
    json_path = pptx_path.with_suffix(".json")
    for f in (pptx_path, json_path):
        if f.exists():
            f.unlink()

    date_str, time_str = pptx_path.stem.split("-")
    specs_leaf = ROOT / "worship" / "specs" / date_str[0:4] / date_str[4:6] / date_str[6:8] / time_str
    if specs_leaf.is_dir():
        for f in specs_leaf.iterdir():
            f.unlink()
        specs_leaf.rmdir()


class TestSongLeaderRequiredField:
    """Regression: on normal (non-custom) templates, the Song Leader field
    could be left blank and the presentation would still generate with no
    leader name anywhere -- it's now required before Generate is enabled.
    Custom templates give every song its own distinct leader field instead
    (see TestCustomTemplateDistinctLeaders), so this requirement is specific
    to the single shared "Song Leader" position on normal templates.
    """

    def test_generate_disabled_until_leader_is_filled(self):
        at = start_normal_template()
        assert _generate_button(at).disabled is True

        at.text_input(key="leader_Song Leader").set_value("Alice").run()
        assert _generate_button(at).disabled is False

    def test_whitespace_only_leader_still_counts_as_missing(self):
        at = start_normal_template()
        at.text_input(key="leader_Song Leader").set_value("   ").run()
        assert _generate_button(at).disabled is True

    def test_custom_template_does_not_require_a_leader(self):
        at = start_custom_template_with_songs(1)
        at.number_input(key="song_song-1").set_value(738).run()
        assert _generate_button(at).disabled is False


class TestWelcomeSlideLeaderNotes:
    """Regression: slides.py's add_welcome() has always written the Song
    Leader's name into the welcome slide's speaker notes when the worship
    spec has a top-level 'leader' key -- but ui.py's own from-scratch spec
    builder (a duplicate of worship.py's generate_json) never set that key,
    so the notes were silently never populated when generating through the
    app, even though the underlying slides.py feature worked.
    """

    def test_generated_welcome_slide_notes_include_leader_name(self):
        at = start_normal_template()
        at.text_input(key="leader_Song Leader").set_value("Alice Test Leader").run()
        assert _generate_button(at).disabled is False

        at.button(key=_generate_button(at).key).click().run(timeout=60)
        assert list(at.exception) == []

        pptx_path = at.session_state["generated_files"]["pptx_path"]
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            notes = prs.slides[0].notes_slide.notes_text_frame.text
            assert notes == "Song Leader: Alice Test Leader"
        finally:
            _cleanup_generated_files(pptx_path)


class TestCustomTemplateWelcomeSlideLeaderNotes:
    """Custom templates give every song its own distinct position ("Song
    Leader 1", "Song Leader 2", ...) instead of sharing one literal "Song
    Leader" key (see TestMakeCustomTemplateItem), so the normal-template
    fix above never matched anything there and the welcome slide's notes
    stayed silently empty. Collect all of them instead, in song order.
    """

    def test_all_song_leaders_appear_in_welcome_slide_notes(self):
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()

        at.selectbox(key="custom_item_type_select").set_value("welcome").run()
        at.button(key="custom_add_item").click().run()
        for _ in range(3):
            at.selectbox(key="custom_item_type_select").set_value("song").run()
            at.button(key="custom_add_item").click().run()

        song_inputs = [ni for ni in at.number_input if ni.key and ni.key.startswith("song_song-")]
        for ni in song_inputs:
            at.number_input(key=ni.key).set_value(738).run()

        leader_inputs = sorted(
            (ti for ti in at.text_input if ti.key and ti.key.startswith("leader_Song Leader")),
            key=lambda t: t.key,
        )
        assert len(leader_inputs) == 3
        for ti, name in zip(leader_inputs, ["Alice", "Bob", "Carol"]):
            at.text_input(key=ti.key).set_value(name).run()

        generate = _generate_button(at)
        assert generate.disabled is False  # no required-leader gate on custom templates
        at.button(key=generate.key).click().run(timeout=60)
        assert list(at.exception) == []

        pptx_path = at.session_state["generated_files"]["pptx_path"]
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            notes = prs.slides[0].notes_slide.notes_text_frame.text
            assert notes == "Song Leader: Alice, Bob, Carol"
        finally:
            _cleanup_generated_files(pptx_path)

    def test_no_leader_key_set_when_no_song_leaders_filled_in(self):
        # A custom template with no songs (or blank leader fields) shouldn't
        # add an empty/garbage 'leader' entry -- the welcome slide should
        # simply get no leader-name notes, same as before this feature.
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("welcome").run()
        at.button(key="custom_add_item").click().run()

        generate = _generate_button(at)
        at.button(key=generate.key).click().run(timeout=60)
        assert list(at.exception) == []

        pptx_path = at.session_state["generated_files"]["pptx_path"]
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            notes = prs.slides[0].notes_slide.notes_text_frame.text
            assert notes == ""
        finally:
            _cleanup_generated_files(pptx_path)


class TestCustomTemplateAnnouncementsTitleItem:
    """New Custom Template Builder option: an "Announcements" title-card
    slide (mirrors Sermon/Lesson/Report -- same quote/background, own text,
    fade to black afterward). Exercises the full UI wiring (dropdown, added
    item, its own leader field, generation) end to end.
    """

    def test_can_add_item_and_it_gets_its_own_leader_field(self):
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("announcements-title").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        leader_inputs = [ti for ti in at.text_input if ti.key and ti.key.startswith("leader_Announcer")]
        assert len(leader_inputs) == 1

    def test_generates_without_leader_being_required(self):
        # Only the normal-template "Song Leader" field is required (see
        # TestSongLeaderRequiredField); this new item type must not trip
        # that same requirement on the custom template path.
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("announcements-title").run()
        at.button(key="custom_add_item").click().run()

        generate = [b for b in at.button if "Generate" in (b.label or "")][0]
        assert generate.disabled is False

        at.button(key=generate.key).click().run(timeout=60)
        assert list(at.exception) == []

        pptx_path = at.session_state["generated_files"]["pptx_path"]
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            assert len(prs.slides) == 2  # title card + fade-to-black
        finally:
            _cleanup_generated_files(pptx_path)


class TestCustomTemplateInvitationSongItems:
    """New Custom Template Builder options: "song-title" (a title-only
    preview slide) and "song-music" (a music-only slide, no title) -- the
    same mechanism sunday-am.json uses for the invitation song: preview it
    before the sermon, then play it as background music after. The two are
    linked by sharing one "id" (see make_custom_template_item and the Add
    Item handler), so picking a song for one applies to both.
    """

    def test_song_title_gets_no_leader_field_and_defaults_its_bubble(self):
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("song-title").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        items = at.session_state["custom_template_items"]
        assert items == [{"type": "song-title", "id": "song-title-1", "bubble": "Invitation Song"}]
        # It previews an already-led song, so it shouldn't ask for its own leader.
        assert not any(ti.key and ti.key.startswith("leader_") for ti in at.text_input)

    def test_song_music_reuses_the_preceding_song_titles_id(self):
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("song-title").run()
        at.button(key="custom_add_item").click().run()
        at.selectbox(key="custom_item_type_select").set_value("song-music").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        items = at.session_state["custom_template_items"]
        assert items[0]["id"] == items[1]["id"] == "song-title-1"
        assert items[1]["type"] == "song-music"
        assert items[1]["fade"] == "out"

    def test_song_music_without_a_song_title_gets_its_own_id(self):
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("song-music").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        items = at.session_state["custom_template_items"]
        assert items == [{"type": "song-music", "id": "song-music-1", "fade": "out"}]

    def test_reordering_a_linked_pair_does_not_collide(self):
        """Regression: the position-reorder box was keyed by item id alone,
        which broke as soon as two items shared an id (a song-title/
        song-music pair does, by design) -- Streamlit raised
        StreamlitDuplicateElementKey the moment both rendered."""
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("song").run()
        at.button(key="custom_add_item").click().run()
        at.selectbox(key="custom_item_type_select").set_value("song-title").run()
        at.button(key="custom_add_item").click().run()
        at.selectbox(key="custom_item_type_select").set_value("song-music").run()
        at.button(key="custom_add_item").click().run()
        assert list(at.exception) == []

        at.number_input(key="custom_pos_song-title-1_song-title").set_value(1).run()
        at.button(key="custom_apply_order").click().run()
        assert list(at.exception) == []

        order = [(it.get("id"), it.get("type")) for it in at.session_state["custom_template_items"]]
        assert order == [("song-title-1", "song-title"), ("song-1", "song"), ("song-title-1", "song-music")]

    def test_full_generation_shares_the_song_between_preview_and_music(self):
        at = AppTest.from_file(str(UI_PY))
        at.run(timeout=30)
        at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()
        at.selectbox(key="custom_item_type_select").set_value("song-title").run()
        at.button(key="custom_add_item").click().run()
        at.selectbox(key="custom_item_type_select").set_value("sermon").run()
        at.button(key="custom_add_item").click().run()
        at.selectbox(key="custom_item_type_select").set_value("song-music").run()
        at.button(key="custom_add_item").click().run()

        # Only the song-music item renders the manual book/number picker --
        # song-title just mirrors it via the shared id at generation time.
        at.number_input(key="song_song-title-1").set_value(738).run()
        leader_inputs = [ti for ti in at.text_input if ti.key and ti.key.startswith("leader_Preach")]
        for ti in leader_inputs:
            at.text_input(key=ti.key).set_value("Alice").run()

        generate = [b for b in at.button if "Generate" in (b.label or "")][0]
        at.button(key=generate.key).click().run(timeout=60)
        assert list(at.exception) == []

        pptx_path = at.session_state["generated_files"]["pptx_path"]
        try:
            spec_path = Path(pptx_path).with_suffix(".json")
            spec = json.loads(spec_path.read_text())
            by_type = {it["type"]: it for it in spec["items"]}
            assert by_type["song-title"]["song"] == "738"
            assert by_type["song-music"]["song"] == "738"
        finally:
            _cleanup_generated_files(pptx_path)


class TestAutofillErrorVisibility:
    """Regression: a failed church-schedule autofill (bad auth, network
    issue, the schedule API genuinely having nothing for a given date) used
    to leave every reading/leader field silently blank with zero feedback --
    indistinguishable from "there's nothing to autofill". A user seeing a
    blank field had no way to tell a real failure from a data gap.
    """

    def test_reading_fetch_error_shows_a_warning(self):
        from unittest.mock import patch
        with patch("worship.fetch_readings", return_value={"readings": {}, "_error": "Scripture API response 401: Unauthorized"}), \
             patch("worship.fetch_leaders", return_value={"leaders": {}}):
            at = AppTest.from_file(str(UI_PY))
            at.run(timeout=30)
            at.selectbox(key="template_select").set_value("sunday-am").run()

        assert list(at.exception) == []
        warnings = [w.value for w in at.warning]
        assert any("Scripture reading" in w and "401: Unauthorized" in w for w in warnings)

    def test_leader_fetch_error_shows_a_warning(self):
        from unittest.mock import patch
        with patch("worship.fetch_leaders", return_value={"leaders": {}, "_error": "Assignments API exception: timed out"}), \
             patch("worship.fetch_readings", return_value={"readings": {}}):
            at = AppTest.from_file(str(UI_PY))
            at.run(timeout=30)
            at.selectbox(key="template_select").set_value("sunday-am").run()

        assert list(at.exception) == []
        warnings = [w.value for w in at.warning]
        assert any("Leader assignments" in w and "timed out" in w for w in warnings)

    def test_no_warning_when_autofill_succeeds(self):
        from unittest.mock import patch
        with patch("worship.fetch_readings", return_value={"readings": {}}), \
             patch("worship.fetch_leaders", return_value={"leaders": {}}):
            at = AppTest.from_file(str(UI_PY))
            at.run(timeout=30)
            at.selectbox(key="template_select").set_value("sunday-am").run()

        assert list(at.exception) == []
        warnings = [w.value for w in at.warning]
        assert not any("Could not auto-fill" in w for w in warnings)

    def test_exception_during_autofill_shows_a_warning_not_a_crash(self):
        from unittest.mock import patch
        with patch("worship.fetch_leaders", side_effect=RuntimeError("boom")):
            at = AppTest.from_file(str(UI_PY))
            at.run(timeout=30)
            at.selectbox(key="template_select").set_value("sunday-am").run()

        assert list(at.exception) == []
        warnings = [w.value for w in at.warning]
        assert any("Could not auto-fill from the church schedule" in w and "boom" in w for w in warnings)


class TestLoadPastService:
    """New "Load Past Service" tab: every Generate click already writes the
    full merged worship JSON (order, types, ids, and each item's own song/
    leader/reading values) to worship/<year>/<date>-<time>.json -- this
    reloads one of those back into the Custom Template Builder instead of
    reverse-engineering the .pptx itself, which is a much bigger and less
    reliable undertaking (some item types are only distinguishable by their
    on-slide label text, not by slide layout).

    Uses a fixture dated 2099 so it always sorts first (most recent) in the
    "Past service" dropdown regardless of what real service files already
    exist in the repo, letting the test rely on the default selection
    instead of needing to pick a specific option out of a selectbox whose
    options are dicts (a combination streamlit.testing.v1.AppTest doesn't
    let a test select by anything other than the default index).
    """

    @staticmethod
    def _write_fixture():
        fixture_dir = ROOT / "worship" / "2099"
        fixture_dir.mkdir(parents=True, exist_ok=True)
        fixture_path = fixture_dir / "20990101-1000.json"
        fixture_path.write_text(json.dumps({
            "isodate": "2099-01-01T10:00:00",
            "template": "custom",
            "type": "Sun - AM",
            "items": [
                {
                    "type": "song", "id": "song-1", "position": "Song Leader 1",
                    "leader": "Alice", "book": "pftl", "song": "738", "coda": 0,
                },
                {"type": "prayer", "id": "prayer-1", "position": "Prayer 1"},
            ],
        }), encoding="utf-8")
        return fixture_dir, fixture_path

    def test_reloads_item_order_song_number_and_leader(self):
        fixture_dir, fixture_path = self._write_fixture()
        try:
            at = AppTest.from_file(str(UI_PY))
            at.run(timeout=30)
            at.selectbox(key="template_select").set_value("Custom Template (Build Order)").run()

            sb = at.selectbox(key="load_past_service_choice")
            assert "2099-01-01" in sb.options[0]

            at.button(key="load_past_service_button").click().run(timeout=30)
            assert list(at.exception) == []

            items = at.session_state["custom_template_items"]
            assert [it["id"] for it in items] == ["song-1", "prayer-1"]
            assert at.number_input(key="song_song-1").value == 738
            assert at.text_input(key="leader_Song Leader 1").value == "Alice"
        finally:
            fixture_path.unlink(missing_ok=True)
            try:
                fixture_dir.rmdir()
            except OSError:
                pass
